"""Positional glyph asymmetry / allograph hypothesis in user-provided IVTFF EVA v3b.
Conservative running-text tokens, each at least 2 characters, no uncertain/editorial matter.
Positions are EVA transcription positions, not validated original character identities.
"""
from docx import Document
import re,collections,json
from pathlib import Path
P=Path('/mnt/data')
rows=[]; langs={}
for p in Document(P/'IVTFF Eva.txt.docx').paragraphs:
 s=p.text.strip()
 if not s:continue
 h=re.match(r'^<(f\d+[rv](?:\d+)?)>\s*<!.*?\$L=([AB])',s)
 if h: langs[h[1]]=h[2]
 m=re.match(r'^<(f\d+[rv](?:\d+)?)\.(\d+),([^>]+)>\s*(.*)',s)
 if not m:continue
 fol,line,loc,src=m.groups()
 if 'P' not in loc:continue
 src=re.sub(r'<![^>]*>',' ',src)
 for word in re.split(r'[.,\s]+',src):
  if re.fullmatch('[a-z]{2,20}',word):rows.append((fol,langs.get(fol,'?'),word))
freq=collections.Counter(w for _,_,w in rows)
counts=collections.Counter(); position=collections.defaultdict(lambda:collections.Counter()); end_context=collections.Counter();
for f,l,w in rows:
 for j,c in enumerate(w):
  loc='initial' if j==0 else 'final' if j==len(w)-1 else 'middle'
  counts[c]+=1;position[c][loc]+=1
 if w.endswith('n'):
  end_context[w[-2]]+=1
line_totals=collections.Counter(); line_last=collections.Counter()
for p in Document(P/'IVTFF Eva.txt.docx').paragraphs:
 s=p.text.strip()
 mm=re.match(r'^<(f\d+[rv](?:\d+)?)\.(\d+),([^>]+)>\s*(.*)',s)
 if not mm or 'P' not in mm[3]: continue
 src=re.sub(r'<![^>]*>',' ',mm[4]);toks=[w for w in re.split(r'[.,\s]+',src) if re.fullmatch('[a-z]{2,20}',w)]
 for j,w in enumerate(toks):
  if w.endswith(('m','n','y')):
   line_totals[w[-1]]+=1
   if j==len(toks)-1:line_last[w[-1]]+=1
out={'method':'Conservative clean EVA running-text tokens P loci; glyph positions are EVA codepoint positions, so multi-codepoint glyphs may be split.', 'tokens':len(rows),'distinct':len(freq),'folios':len({f for f,_,_ in rows}), 'glyph_positions':{},'n_final_preceded_by':end_context.most_common(15), 'line_final_among_word_final':{c:{'at_last_parsed_token':line_last[c],'all':line_totals[c],'pct':round(line_last[c]*100/line_totals[c],2)} for c in 'mny'}}
for c in ['i','n','m','q','o','a','y','l','r','d','e','k','t','p','f','s','h']:
 cc=position[c];total=sum(cc.values())
 out['glyph_positions'][c]={'initial':cc['initial'],'middle':cc['middle'],'final':cc['final'],'total':total,'pct_final':round(cc['final']*100/total,2) if total else 0,'pct_initial':round(cc['initial']*100/total,2) if total else 0}
# special comparative facts across variations
out['words_ending_i_examples']=freq.most_common(2000) # trimmed below
out['words_ending_i_examples']=[(w,n) for w,n in freq.most_common() if w.endswith('i')][:20]
out['words_ending_n_examples']=[(w,n) for w,n in freq.most_common() if w.endswith('n')][:20]
out['words_ending_m_examples']=[(w,n) for w,n in freq.most_common() if w.endswith('m')][:20]
out['common_m_medial_examples']=[(w,n) for w,n in freq.most_common() if 'm' in w[:-1]][:20]
# test paired forms where final 'i' versus 'n' occurs for same prefix; exact minimal pairs
pairs=[]
for w,n in freq.items():
 if w.endswith('n') and w[:-1]+'i' in freq:
  pairs.append({'stem':w[:-1],'n_form':w,'n_freq':n,'i_form':w[:-1]+'i','i_freq':freq[w[:-1]+'i']})
out['final_n_vs_i_minimal_pairs']=sorted(pairs,key=lambda x:-x['n_freq'])[:30]
out['num_final_n_vs_i_minimal_pairs']=len(pairs)
(P/'voynich_experiment_051_results.json').write_text(json.dumps(out,indent=2))
print('tokens',len(rows),'distinct',len(freq),'folios',out['folios'])
for c in ['i','n','m','q','o','a','y','l','r','d','e','k','t']:
 print(c,out['glyph_positions'][c])
print('n prefinal',out['n_final_preceded_by'][:8]);print('i final examples',out['words_ending_i_examples'][:10]);print('minimal pairs',pairs[:10]);print('m medial',out['common_m_medial_examples'][:10])
